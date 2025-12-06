#!/usr/bin/env python3
"""
Create Word Report with Minimalistic Academic Style
"""

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
import os

# Paths
BASE_DIR = "/Users/yiwei/GithubRepos/GameData_Analysis_Projects/Delivery 2 A_B_testing"
VIZ_DIR = os.path.join(BASE_DIR, "Analysis/Visualization")
OUTPUT_PATH = os.path.join(BASE_DIR, "Deliverables/Report_Final.docx")

# Create document
doc = Document()

# Set up styles
style = doc.styles['Normal']
font = style.font
font.name = 'Calibri'
font.size = Pt(11)

# Title style
title_style = doc.styles['Title']
title_style.font.size = Pt(24)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor(0, 0, 0)

# Heading styles
for i in range(1, 4):
    h_style = doc.styles[f'Heading {i}']
    h_style.font.name = 'Calibri'
    h_style.font.color.rgb = RGBColor(0, 0, 0)
    h_style.font.bold = True
    if i == 1:
        h_style.font.size = Pt(16)
    elif i == 2:
        h_style.font.size = Pt(14)
    else:
        h_style.font.size = Pt(12)

# =============================================================================
# TITLE PAGE
# =============================================================================
doc.add_paragraph()
doc.add_paragraph()
doc.add_paragraph()

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("A/B Test Analysis Report")
run.font.size = Pt(28)
run.font.bold = True

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run("Boost Power Evaluation")
run.font.size = Pt(18)
run.font.color.rgb = RGBColor(100, 100, 100)

doc.add_paragraph()
doc.add_paragraph()

info = doc.add_paragraph()
info.alignment = WD_ALIGN_PARAGRAPH.CENTER
info.add_run("Game Data Analysis\n").font.size = Pt(12)
info.add_run("Delivery 2 - A/B Testing\n").font.size = Pt(12)
info.add_run("December 2024").font.size = Pt(12)

doc.add_page_break()

# =============================================================================
# EXECUTIVE SUMMARY
# =============================================================================
doc.add_heading('1. Executive Summary', level=1)

doc.add_paragraph(
    "This report analyzes a 4-month A/B test conducted on a puzzle game to determine "
    "whether an increased boost power feature should be rolled out to all players."
)

doc.add_heading('Key Findings', level=3)
bullets = [
    "Players with stronger boost (Group A) achieved a 2.6 percentage point higher success rate (25.1% vs 22.5%)",
    "The difference is statistically significant (p < 0.001)",
    "Group A players used boosts 51% more frequently than Group B (15.6% vs 10.3%)",
    "The positive effect persisted even after the boost was reset to original power"
]
for bullet in bullets:
    doc.add_paragraph(bullet, style='List Bullet')

doc.add_heading('Recommendation', level=3)
doc.add_paragraph(
    "Roll out the stronger boost to all players, with ongoing monitoring of "
    "long-term retention and monetization metrics."
).bold = True

# =============================================================================
# INTRODUCTION
# =============================================================================
doc.add_heading('2. Introduction & Background', level=1)

doc.add_heading('2.1 Game Context', level=2)
doc.add_paragraph(
    "The game is a Candy Crush-style puzzle game where players attempt to complete levels. "
    "Players can use 'boosts' - power-ups that help them succeed at difficult levels. "
    "The game launched on January 1, 2020."
)

doc.add_heading('2.2 Business Problem', level=2)
doc.add_paragraph(
    "Players reported that some levels were too difficult, leading to frustration. "
    "The product team hypothesized that increasing the boost power would improve player "
    "experience and engagement without negatively impacting game balance."
)

doc.add_heading('2.3 Test Design', level=2)

# Test design table
table = doc.add_table(rows=4, cols=2)
table.style = 'Table Grid'
table.alignment = WD_TABLE_ALIGNMENT.CENTER

cells = [
    ("Group A (Test)", "Received increased boost power"),
    ("Group B (Control)", "Retained original boost power"),
    ("Test Period", "June 1 - September 30, 2020"),
    ("Sample Size", "~14,000 players (6,923 A, 7,017 B)")
]

for i, (label, value) in enumerate(cells):
    table.rows[i].cells[0].text = label
    table.rows[i].cells[1].text = value
    table.rows[i].cells[0].paragraphs[0].runs[0].bold = True

doc.add_paragraph()

doc.add_heading('2.4 Research Question', level=2)
p = doc.add_paragraph()
p.add_run("Should the increased boost power be rolled out to all players?").italic = True

# =============================================================================
# METHODOLOGY
# =============================================================================
doc.add_heading('3. Methodology', level=1)

doc.add_heading('3.1 Data Sources', level=2)

table = doc.add_table(rows=5, cols=3)
table.style = 'Table Grid'
headers = ["Dataset", "Records", "Key Fields"]
for i, h in enumerate(headers):
    table.rows[0].cells[i].text = h
    table.rows[0].cells[i].paragraphs[0].runs[0].bold = True

data = [
    ("Users", "13,940", "TestGroup, demographics"),
    ("Sessions", "432,000+", "Session timing"),
    ("Levels", "1.5M+", "Attempts, boost, outcomes"),
    ("Transactions", "26,000+", "Purchase data")
]
for i, row in enumerate(data):
    for j, val in enumerate(row):
        table.rows[i+1].cells[j].text = val

doc.add_paragraph()

doc.add_heading('3.2 Time Period Segmentation', level=2)

table = doc.add_table(rows=4, cols=3)
table.style = 'Table Grid'
headers = ["Period", "Dates", "Purpose"]
for i, h in enumerate(headers):
    table.rows[0].cells[i].text = h
    table.rows[0].cells[i].paragraphs[0].runs[0].bold = True

periods = [
    ("Pre-Test", "Jan 1 - May 31, 2020", "Baseline validation"),
    ("During Test", "June 1 - Sept 30, 2020", "A/B comparison"),
    ("Post-Test", "Oct 1 - Dec 31, 2020", "Persistence analysis")
]
for i, row in enumerate(periods):
    for j, val in enumerate(row):
        table.rows[i+1].cells[j].text = val

doc.add_paragraph()

doc.add_heading('3.3 Statistical Methods', level=2)
methods = [
    "Chi-squared test for comparing proportions between groups",
    "95% Confidence Intervals for effect size estimation",
    "Odds Ratio for practical significance interpretation",
    "Significance threshold: α = 0.05"
]
for m in methods:
    doc.add_paragraph(m, style='List Bullet')

# =============================================================================
# RESULTS
# =============================================================================
doc.add_heading('4. Results & Analysis', level=1)

doc.add_heading('4.1 Baseline Validation', level=2)
doc.add_paragraph(
    "Before analyzing test results, we verified that Groups A and B were comparable "
    "before the test began."
)

table = doc.add_table(rows=3, cols=4)
table.style = 'Table Grid'
headers = ["Metric", "Group A", "Group B", "p-value"]
for i, h in enumerate(headers):
    table.rows[0].cells[i].text = h
    table.rows[0].cells[i].paragraphs[0].runs[0].bold = True

baseline_data = [
    ("Sample Size", "6,923", "7,017", "-"),
    ("Pre-Test Success Rate", "29.1%", "30.1%", "0.0016")
]
for i, row in enumerate(baseline_data):
    for j, val in enumerate(row):
        table.rows[i+1].cells[j].text = val

doc.add_paragraph()
doc.add_paragraph(
    "Finding: There was a small but statistically significant difference in baseline "
    "success rates (1 percentage point). This is noted as a limitation, but the "
    "difference is small enough that it does not invalidate the test results."
)

# PRIMARY FINDING
doc.add_heading('4.2 Primary Finding: Level Success Rate', level=2)

doc.add_paragraph("During the test period (June - September 2020):")

# Add visualization
viz_path = os.path.join(VIZ_DIR, "Success Rate by Group.png")
if os.path.exists(viz_path):
    doc.add_picture(viz_path, width=Inches(5))
    last_paragraph = doc.paragraphs[-1]
    last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption = doc.add_paragraph("Figure 1: Success Rate by Test Group")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.runs[0].italic = True
    caption.runs[0].font.size = Pt(10)

doc.add_paragraph()

table = doc.add_table(rows=3, cols=4)
table.style = 'Table Grid'
headers = ["Metric", "Group A", "Group B", "Difference"]
for i, h in enumerate(headers):
    table.rows[0].cells[i].text = h
    table.rows[0].cells[i].paragraphs[0].runs[0].bold = True

table.rows[1].cells[0].text = "Success Rate"
table.rows[1].cells[1].text = "25.1%"
table.rows[1].cells[2].text = "22.5%"
table.rows[1].cells[3].text = "+2.6 pp"
table.rows[2].cells[0].text = "Level Attempts"
table.rows[2].cells[1].text = "399,219"
table.rows[2].cells[2].text = "394,339"
table.rows[2].cells[3].text = "-"

doc.add_paragraph()
doc.add_heading('Statistical Test Results', level=3)

stats = [
    "Chi-squared statistic: 797.84",
    "Degrees of freedom: 1",
    "p-value: < 2.2e-16 (highly significant)",
    "Odds Ratio: 1.15",
    "95% CI for difference: [2.35%, 2.86%]"
]
for s in stats:
    doc.add_paragraph(s, style='List Bullet')

# TIME SERIES
doc.add_heading('4.3 Time-Series Analysis', level=2)

viz_path = os.path.join(VIZ_DIR, "Success Rate Over Time.png")
if os.path.exists(viz_path):
    doc.add_picture(viz_path, width=Inches(5.5))
    last_paragraph = doc.paragraphs[-1]
    last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption = doc.add_paragraph("Figure 2: Success Rate Over Time by Test Group")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.runs[0].italic = True
    caption.runs[0].font.size = Pt(10)

doc.add_paragraph()
doc.add_paragraph(
    "The time-series visualization shows clear separation during the test period, "
    "with Group A consistently outperforming Group B. Notably, the effect persists "
    "even after the boost was reset in October."
)

# BOOST ENGAGEMENT
doc.add_heading('4.4 Secondary Finding: Boost Engagement', level=2)

viz_path = os.path.join(VIZ_DIR, "Boost Usage Rate.png")
if os.path.exists(viz_path):
    doc.add_picture(viz_path, width=Inches(5))
    last_paragraph = doc.paragraphs[-1]
    last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption = doc.add_paragraph("Figure 3: Boost Usage Rate by Test Group")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.runs[0].italic = True
    caption.runs[0].font.size = Pt(10)

doc.add_paragraph()

table = doc.add_table(rows=2, cols=4)
table.style = 'Table Grid'
headers = ["Metric", "Group A", "Group B", "Difference"]
for i, h in enumerate(headers):
    table.rows[0].cells[i].text = h
    table.rows[0].cells[i].paragraphs[0].runs[0].bold = True

table.rows[1].cells[0].text = "Boost Usage Rate"
table.rows[1].cells[1].text = "15.6%"
table.rows[1].cells[2].text = "10.3%"
table.rows[1].cells[3].text = "+51%"

doc.add_paragraph()
doc.add_paragraph(
    "Players used the boost significantly more when it was stronger (p < 0.001), "
    "indicating they recognized and appreciated the improved feature value."
)

# PERSISTENCE
doc.add_heading('4.5 Persistence Effect', level=2)

viz_path = os.path.join(VIZ_DIR, "Success Rate by Test Period.png")
if os.path.exists(viz_path):
    doc.add_picture(viz_path, width=Inches(5.5))
    last_paragraph = doc.paragraphs[-1]
    last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption = doc.add_paragraph("Figure 4: Success Rate by Test Period")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.runs[0].italic = True
    caption.runs[0].font.size = Pt(10)

doc.add_paragraph()

table = doc.add_table(rows=4, cols=4)
table.style = 'Table Grid'
headers = ["Period", "Group A", "Group B", "Difference"]
for i, h in enumerate(headers):
    table.rows[0].cells[i].text = h
    table.rows[0].cells[i].paragraphs[0].runs[0].bold = True

period_data = [
    ("Pre-Test", "29.1%", "30.1%", "-1.0 pp"),
    ("During Test", "25.1%", "22.5%", "+2.6 pp"),
    ("Post-Test", "22.3%", "19.7%", "+2.6 pp")
]
for i, row in enumerate(period_data):
    for j, val in enumerate(row):
        table.rows[i+1].cells[j].text = val

doc.add_paragraph()
doc.add_paragraph(
    "Key Insight: Even after the boost was reset to original power, Group A continued "
    "to outperform Group B by the same margin. This suggests learned behavior or "
    "habit formation - players who experienced the stronger boost developed better "
    "strategies or confidence that persisted."
).bold = True

# =============================================================================
# DISCUSSION
# =============================================================================
doc.add_heading('5. Discussion', level=1)

doc.add_heading('5.1 Arguments FOR Rolling Out Stronger Boost', level=2)
for_args = [
    "Significant Performance Improvement: 2.6 percentage point increase with p < 0.001",
    "Higher Player Engagement: 51% increase in boost usage indicates feature value",
    "Lasting Positive Effects: Effect persists after reset, suggesting skill development",
    "No Observed Negative Effects: No decline in session frequency or engagement"
]
for arg in for_args:
    doc.add_paragraph(arg, style='List Bullet')

doc.add_heading('5.2 Arguments AGAINST Rolling Out Stronger Boost', level=2)
against_args = [
    "May Reduce Game Challenge: Easier success could reduce long-term engagement",
    "Potential Boost Dependency: Players may become reliant on boosts",
    "Unknown Monetization Impact: Stronger free boost might reduce purchases",
    "Baseline Difference: Small pre-test difference suggests imperfect randomization"
]
for arg in against_args:
    doc.add_paragraph(arg, style='List Bullet')

# =============================================================================
# CONCLUSION
# =============================================================================
doc.add_heading('6. Conclusion & Recommendation', level=1)

doc.add_heading('6.1 Summary of Findings', level=2)

table = doc.add_table(rows=5, cols=3)
table.style = 'Table Grid'
headers = ["Test", "Result", "Significant?"]
for i, h in enumerate(headers):
    table.rows[0].cells[i].text = h
    table.rows[0].cells[i].paragraphs[0].runs[0].bold = True

summary_data = [
    ("Success Rate (During)", "A > B by 2.6 pp", "Yes (p < 0.001)"),
    ("Boost Usage (During)", "A > B by 51%", "Yes (p < 0.001)"),
    ("Baseline (Pre-Test)", "Similar (1 pp diff)", "Minor concern"),
    ("Persistence (Post-Test)", "Effect persists", "Yes (p < 0.001)")
]
for i, row in enumerate(summary_data):
    for j, val in enumerate(row):
        table.rows[i+1].cells[j].text = val

doc.add_paragraph()

doc.add_heading('6.2 Final Recommendation', level=2)
p = doc.add_paragraph()
run = p.add_run("RECOMMEND: Roll out the stronger boost to all players")
run.bold = True
run.font.size = Pt(12)

doc.add_paragraph()
doc.add_paragraph("The evidence strongly supports this decision:")
evidence = [
    "Clear, statistically significant improvement in player success",
    "Increased engagement with the boost feature",
    "Positive lasting effects on player behavior",
    "No observed negative consequences during the test period"
]
for e in evidence:
    doc.add_paragraph(e, style='List Bullet')

doc.add_heading('6.3 Suggested Next Steps', level=2)
next_steps = [
    "Monitor Long-Term Metrics: Track D7, D30, D90 retention after rollout",
    "Analyze Monetization Impact: Compare ARPU and boost purchase behavior",
    "Consider Level Rebalancing: Review difficulty curves if needed",
    "A/B Test Monetization: Test pricing changes if boost purchases decline"
]
for step in next_steps:
    doc.add_paragraph(step, style='List Bullet')

# Save document
doc.save(OUTPUT_PATH)
print(f"Report saved to: {OUTPUT_PATH}")
